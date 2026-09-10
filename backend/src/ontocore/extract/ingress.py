from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from ontocore.errors import IngressError
from ontocore.models import ParsedDocument, TextBlock

_PERIOD_RE = re.compile(r"[。.]")


_TXT_ENCODINGS = ("utf-8", "gb18030")


def parse_upload(filename: str, data: bytes) -> ParsedDocument:
    if not data:
        raise IngressError("无法提取文本")

    suffix = Path(filename).suffix.lower()
    if suffix == ".txt":
        text = _decode_txt(data)
        return _build_document(filename, text)
    if suffix == ".pdf":
        text = _extract_pdf(data)
        return _build_document(filename, text)
    if suffix == ".docx":
        text = _extract_docx(data)
        return _build_document(filename, text)
    raise IngressError()


def _decode_txt(data: bytes) -> str:
    for encoding in _TXT_ENCODINGS:
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise IngressError("无法提取文本")


def _extract_pdf(data: bytes) -> str:
    try:
        reader = PdfReader(BytesIO(data))
        if reader.is_encrypted:
            raise IngressError("无法提取文本")
        parts: list[str] = []
        for page in reader.pages:
            parts.append(page.extract_text() or "")
        text = "\n".join(parts)
        if not text.strip():
            raise IngressError("无法提取文本")
        return text
    except IngressError:
        raise
    except Exception:
        raise IngressError("无法提取文本") from None


def _extract_docx(data: bytes) -> str:
    try:
        document = Document(BytesIO(data))
        paragraphs = [
            paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()
        ]
        if not paragraphs:
            raise IngressError("无法提取文本")
        return "\n\n".join(paragraphs)
    except IngressError:
        raise
    except Exception:
        raise IngressError("无法提取文本") from None


def _build_document(filename: str, text: str) -> ParsedDocument:
    chunks = _split_chunks(text)
    blocks = tuple(
        TextBlock(f"b{index}", _block_kind(chunk), _normalize_chunk(chunk))
        for index, chunk in enumerate(chunks)
    )
    return ParsedDocument(filename, text, blocks)


def _split_chunks(text: str) -> list[str]:
    chunks = [chunk.strip() for chunk in re.split(r"\n\s*\n", text) if chunk.strip()]
    if not chunks:
        raise IngressError("无法提取文本")
    return chunks


def _block_kind(chunk: str) -> str:
    lines = [line.strip() for line in chunk.splitlines() if line.strip()]
    if len(lines) != 1:
        return "paragraph"
    line = lines[0]
    if line.startswith("#"):
        return "heading"
    if len(line) < 40 and not _PERIOD_RE.search(line):
        return "heading"
    return "paragraph"


def _normalize_chunk(chunk: str) -> str:
    lines = [line.strip() for line in chunk.splitlines() if line.strip()]
    if len(lines) == 1 and lines[0].startswith("#"):
        return lines[0].lstrip("#").strip()
    return chunk.strip()
